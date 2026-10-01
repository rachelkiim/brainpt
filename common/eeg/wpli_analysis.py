"""Frontoparietal wPLI connectivity (pre-specified channel pairs)."""
import numpy as np, mne, os, re, pickle, sys
import openpyxl
mne.set_log_level('ERROR')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from eeg_pipeline import run_preprocessing, EEG_CHANNELS

PAIRS = [('F3','P7'), ('F4','P8'), ('AF3','P7'), ('AF4','P8')]
BANDS = {'Theta':(4,8), 'Alpha':(8,12), 'Beta':(12,30)}

def wpli(x, y, fs, band, epoch_sec=2.0, overlap=0.5):
    """
    weighted Phase Lag Index (Vinck et al. 2011, NeuroImage)
    wPLI = |E{Im(Sxy)}| / E{|Im(Sxy)|}, expectation over epochs.
    """
    n = int(fs*epoch_sec)
    step = int(n*(1-overlap))
    if len(x) < n*3:   # at least 3 epochs
        return np.nan
    win = np.hanning(n)
    starts = range(0, len(x)-n+1, step)
    Sxy = []
    for s0 in starts:
        xs = (x[s0:s0+n] - x[s0:s0+n].mean()) * win
        ys = (y[s0:s0+n] - y[s0:s0+n].mean()) * win
        X = np.fft.rfft(xs); Y = np.fft.rfft(ys)
        Sxy.append(X * np.conj(Y))
    Sxy = np.array(Sxy)                      # (n_epoch, n_freq)
    if Sxy.shape[0] < 3: return np.nan
    freqs = np.fft.rfftfreq(n, 1/fs)
    idx = (freqs>=band[0]) & (freqs<=band[1])
    if not idx.any(): return np.nan
    im = np.imag(Sxy[:, idx])                # (n_epoch, n_band_freq)
    num = np.abs(im.mean(axis=0))            # mean over epochs, per frequency
    den = np.abs(im).mean(axis=0)
    with np.errstate(divide='ignore', invalid='ignore'):
        w = np.where(den>0, num/den, np.nan)
    return float(np.nanmean(w))              # mean over frequencies in band


def load_marker(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    return [(float(r[0]), float(r[1]), str(r[2])) for r in rows[1:] if r[0] is not None and r[2] is not None]

def process(edf, marker_path, trim_s=30, trim_e=20):
    rows = load_marker(marker_path)
    cleaned, orig_n = run_preprocessing(edf)
    fs = cleaned.info['sfreq']; scale = cleaned.n_times/orig_n
    ev = {}
    for lat,dur,typ in rows:
        s=max(0,round(lat*scale)); e=min(cleaned.n_times, round((lat+dur)*scale))
        ev[typ]=(s,e)
    if 'baseline' in ev:
        s,e = ev['baseline']; ev['baseline']=(s+int(trim_s*fs), e-int(trim_e*fs))
    data = cleaned.get_data()
    out = {}
    for cond,(s,e) in ev.items():
        if (e-s)/fs < 15: continue
        seg = data[:, s:e]
        for (a,b) in PAIRS:
            ia, ib = EEG_CHANNELS.index(a), EEG_CHANNELS.index(b)
            for bn, br in BANDS.items():
                out[(cond, f"{a}-{b}", bn)] = wpli(seg[ia], seg[ib], fs, br)
                out[(cond, f"PSI_{a}-{b}", bn)] = psi(seg[ia], seg[ib], fs, br)
    return out


def psi(x, y, fs, band, epoch_sec=2.0, overlap=0.5):
    """
    Phase Slope Index (Nolte et al. 2008, PRL)
    Positive = x leads y; negative = y leads x.
    """
    n = int(fs*epoch_sec); step = int(n*(1-overlap))
    if len(x) < n*3: return np.nan
    win = np.hanning(n)
    Sxy=[]; Sxx=[]; Syy=[]
    for s0 in range(0, len(x)-n+1, step):
        xs=(x[s0:s0+n]-x[s0:s0+n].mean())*win
        ys=(y[s0:s0+n]-y[s0:s0+n].mean())*win
        X=np.fft.rfft(xs); Y=np.fft.rfft(ys)
        Sxy.append(X*np.conj(Y)); Sxx.append(np.abs(X)**2); Syy.append(np.abs(Y)**2)
    Sxy=np.array(Sxy).mean(axis=0); Sxx=np.array(Sxx).mean(axis=0); Syy=np.array(Syy).mean(axis=0)
    C = Sxy/np.sqrt(Sxx*Syy)          # complex coherency
    freqs=np.fft.rfftfreq(n,1/fs)
    idx=np.where((freqs>=band[0]) & (freqs<=band[1]))[0]
    if len(idx)<2: return np.nan
    # PSI = Im( sum_f conj(C(f)) * C(f+df) )
    val = np.sum(np.conj(C[idx[:-1]]) * C[idx[1:]])
    return float(np.imag(val))
